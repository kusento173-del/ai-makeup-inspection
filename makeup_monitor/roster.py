from __future__ import annotations

import logging
import shutil
from dataclasses import dataclass, replace
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .config import resolve_path


LOGGER = logging.getLogger(__name__)


@dataclass(frozen=True)
class Assignment:
    anchor_code: str
    streamer_name: str
    makeup_artist: str
    phone: str
    team: str
    start_date: date | None
    end_date: date | None


@dataclass(frozen=True)
class RosterEntry:
    anchor_code: str
    streamer_name: str
    makeup_artist: str
    phone: str
    douyin_accounts: tuple[str, ...]
    team: str = ""


@dataclass(frozen=True)
class MakeupArtistContact:
    phone: str
    team: str


@dataclass(frozen=True)
class RosterSnapshot:
    roster_file: Path
    raw_file: Path
    entries: tuple[RosterEntry, ...]
    account_index: dict[str, RosterEntry]
    lookback_start: date
    lookback_end: date


def normalize_identifier(value: Any) -> str:
    if value in (None, "", 0):
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else format(value, "f").rstrip("0").rstrip(".")
    return str(value).strip()


def cell_date(value: Any, epoch: datetime) -> date | None:
    if value in (None, "", 0):
        return None
    if isinstance(value, time) and value == time.min:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch)
        return converted.date() if isinstance(converted, datetime) else converted
    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError(f"无法识别日期: {value!r}")


def latest_backup(backup_dir: Path, patterns: Iterable[str], label: str) -> Path:
    files: list[Path] = []
    for pattern in patterns:
        files.extend(path for path in backup_dir.glob(pattern) if path.is_file())
    if not files:
        joined = "、".join(patterns)
        raise FileNotFoundError(f"没有找到{label}的 WPS 本地备份，目录: {backup_dir}，匹配: {joined}")
    return max(files, key=lambda path: path.stat().st_mtime)


def cached_backup(
    backup_dir: Path,
    patterns: Iterable[str],
    label: str,
    cache_path: Path,
) -> Path:
    source = latest_backup(backup_dir, patterns, label)
    return cache_source(source, label, cache_path)


def cache_source(source: Path, label: str, cache_path: Path) -> Path:
    if cache_path.is_file() and cache_path.stat().st_mtime >= source.stat().st_mtime:
        return cache_path

    cache_path.parent.mkdir(parents=True, exist_ok=True)
    temp = cache_path.with_suffix(f"{cache_path.suffix}.tmp")
    try:
        shutil.copy2(source, temp)
        temp.replace(cache_path)
    except OSError:
        temp.unlink(missing_ok=True)
        if cache_path.is_file():
            LOGGER.warning("%s 的 WPS 备份正在变化，本轮沿用本地缓存。", label)
            return cache_path
        raise
    return cache_path


def configured_source(
    value: Any,
    backup_dir: Path,
    patterns: Iterable[str],
    label: str,
    cache_path: Path,
) -> Path:
    configured = str(value or "").strip()
    if configured:
        direct_path = resolve_path(configured)
        if direct_path.is_file():
            return cache_source(direct_path, label, cache_path)
        LOGGER.warning("%s 的 WPS 同步文件不存在，改用自动恢复备份: %s", label, direct_path)
    return cached_backup(backup_dir, patterns, label, cache_path)


def select_active_assignments(
    rows: Iterable[tuple[Any, ...]],
    epoch: datetime,
    today: date,
    columns: dict[str, int | None] | None = None,
) -> dict[str, Assignment]:
    columns = columns or {
        "anchor_code": 0,
        "streamer_name": 1,
        "makeup_artist": 2,
        "phone": None,
        "start_date": 3,
        "end_date": 4,
        "status": 9,
    }

    def value_at(row: tuple[Any, ...], field: str) -> Any:
        index = columns.get(field)
        return row[index] if index is not None and index < len(row) else None

    assignments: dict[str, Assignment] = {}
    for row_number, row in enumerate(rows, start=2):
        row = tuple(row)
        anchor_code = normalize_identifier(value_at(row, "anchor_code"))
        makeup_artist = normalize_identifier(value_at(row, "makeup_artist"))
        end_date = cell_date(value_at(row, "end_date"), epoch)
        status = normalize_identifier(value_at(row, "status"))
        if not anchor_code or makeup_artist == "未分组" or status != "线下":
            continue
        if end_date and end_date < today:
            continue
        assignment = Assignment(
            anchor_code=anchor_code,
            streamer_name=normalize_identifier(value_at(row, "streamer_name")),
            makeup_artist=makeup_artist,
            phone=normalize_identifier(value_at(row, "phone")),
            team="",
            start_date=cell_date(value_at(row, "start_date"), epoch),
            end_date=end_date,
        )
        previous = assignments.get(anchor_code)
        if previous and previous != assignment:
            raise ValueError(
                f"主播编号 {anchor_code} 存在多个当前有效归属，"
                f"请检查固定主播名单第 {row_number} 行。"
            )
        assignments[anchor_code] = assignment
    return assignments


def assignment_columns(headers: tuple[Any, ...]) -> dict[str, int | None]:
    normalized = {
        normalize_identifier(header): index
        for index, header in enumerate(headers)
        if normalize_identifier(header)
    }
    required = {
        "anchor_code": "主播编号",
        "streamer_name": "主播昵称",
        "makeup_artist": "化妆师",
        "start_date": "开始日期",
        "end_date": "结束日期",
        "status": "状态",
    }
    missing = [header for header in required.values() if header not in normalized]
    if missing:
        raise ValueError(f"固定主播名单缺少表头: {'、'.join(missing)}")
    return {
        field: normalized[header]
        for field, header in required.items()
    } | {"phone": None}


def makeup_staff_columns(headers: tuple[Any, ...]) -> dict[str, int | None]:
    normalized = {
        normalize_identifier(header): index
        for index, header in enumerate(headers)
        if normalize_identifier(header)
    }
    if "化妆师" not in normalized:
        raise ValueError("化妆师人员表缺少“化妆师”表头。")
    if "小队" not in normalized:
        raise ValueError("化妆师人员表缺少“小队”表头。")
    phone = next(
        (
            normalized[name]
            for name in ("手机号", "手机号码", "联系电话", "电话")
            if name in normalized
        ),
        None,
    )
    return {
        "makeup_artist": normalized["化妆师"],
        "phone": phone,
        "team": normalized["小队"],
    }


def select_makeup_artist_contacts(
    rows: Iterable[tuple[Any, ...]],
    columns: dict[str, int | None],
) -> dict[str, MakeupArtistContact]:
    artist_index = columns["makeup_artist"]
    phone_index = columns["phone"]
    team_index = columns["team"]
    assert artist_index is not None
    assert team_index is not None
    contacts: dict[str, MakeupArtistContact] = {}
    for row_number, row in enumerate(rows, start=2):
        row = tuple(row)
        artist = normalize_identifier(row[artist_index] if artist_index < len(row) else None)
        phone = normalize_identifier(
            row[phone_index]
            if phone_index is not None and phone_index < len(row)
            else None
        )
        team = normalize_identifier(row[team_index] if team_index < len(row) else None)
        if team == "/":
            team = ""
        if not artist:
            continue
        previous = contacts.get(artist, MakeupArtistContact("", ""))
        if previous.phone and phone and previous.phone != phone:
            raise ValueError(
                f"化妆师“{artist}”存在多个不同手机号，"
                f"请检查化妆师人员表第 {row_number} 行。"
            )
        if previous.team and team and previous.team != team:
            raise ValueError(
                f"化妆师“{artist}”存在多个不同小队，"
                f"请检查化妆师人员表第 {row_number} 行。"
            )
        contacts[artist] = MakeupArtistContact(
            phone=phone or previous.phone,
            team=team or previous.team,
        )
    return contacts


def load_assignments(
    path: Path,
    sheet_name: str,
    staff_sheet_name: str,
    today: date,
) -> dict[str, Assignment]:
    with path.open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"固定名单缺少工作表: {sheet_name}")
            if staff_sheet_name not in workbook.sheetnames:
                raise ValueError(f"固定名单缺少工作表: {staff_sheet_name}")
            sheet = workbook[sheet_name]
            headers = tuple(
                next(
                    sheet.iter_rows(
                        min_row=1,
                        max_row=1,
                        max_col=100,
                        values_only=True,
                    )
                )
            )
            columns = assignment_columns(headers)
            max_column = max(
                index for index in columns.values() if index is not None
            ) + 1
            assignments = select_active_assignments(
                sheet.iter_rows(min_row=2, max_col=max_column, values_only=True),
                workbook.epoch,
                today,
                columns,
            )

            staff_sheet = workbook[staff_sheet_name]
            staff_headers = tuple(
                next(
                    staff_sheet.iter_rows(
                        min_row=1,
                        max_row=1,
                        max_col=100,
                        values_only=True,
                    )
                )
            )
            staff_columns = makeup_staff_columns(staff_headers)
            staff_max_column = max(
                index for index in staff_columns.values() if index is not None
            ) + 1
            contacts = select_makeup_artist_contacts(
                staff_sheet.iter_rows(
                    min_row=2,
                    max_col=staff_max_column,
                    values_only=True,
                ),
                staff_columns,
            )
            if staff_columns["phone"] is None:
                LOGGER.warning(
                    "%s 暂未找到手机号表头，本轮所有化妆师均不 @。",
                    staff_sheet_name,
                )
            resolved: dict[str, Assignment] = {}
            for anchor_code, assignment in assignments.items():
                contact = contacts.get(
                    assignment.makeup_artist,
                    MakeupArtistContact("", ""),
                )
                resolved[anchor_code] = replace(
                    assignment,
                    phone=contact.phone,
                    team=contact.team,
                )
            return resolved
        finally:
            workbook.close()


def load_recent_douyin_accounts(
    path: Path,
    sheet_name: str,
    assignments: dict[str, Assignment],
    start: date,
    end: date,
) -> dict[str, set[str]]:
    accounts: dict[str, set[str]] = {}
    with path.open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"经纪人原始数据缺少工作表: {sheet_name}")
            sheet = workbook[sheet_name]
            for row in sheet.iter_rows(min_row=2, max_col=18, values_only=True):
                if normalize_identifier(row[0]) != "抖音":
                    continue
                anchor_code = normalize_identifier(row[17])
                if anchor_code not in assignments:
                    continue
                row_date = cell_date(row[1], workbook.epoch)
                if not row_date or row_date < start or row_date > end:
                    continue
                account = normalize_identifier(row[2])
                if account:
                    accounts.setdefault(anchor_code, set()).add(account)
        finally:
            workbook.close()
    return accounts


def load_roster(config: dict[str, Any], today: date | None = None) -> RosterSnapshot:
    today = today or date.today()
    wps = config["wps"]
    backup_dir = resolve_path(wps["backup_dir"])
    cache_dir = resolve_path(config.get("data_root", "daily_data")) / "_source_cache"
    roster_file = configured_source(
        wps.get("roster_file"),
        backup_dir,
        wps["roster_backup_globs"],
        "固定主播名单",
        cache_dir / "fixed_roster.xlsx",
    )
    raw_file = configured_source(
        wps.get("raw_data_file"),
        backup_dir,
        wps["raw_data_backup_globs"],
        "2026年经纪人原始数据",
        cache_dir / "raw_data.xlsx",
    )

    lookback_days = max(1, int(wps.get("lookback_days", 60)))
    lookback_start = today - timedelta(days=lookback_days - 1)
    assignments = load_assignments(
        roster_file,
        wps["roster_sheet"],
        wps.get("staff_sheet", "化妆师人员表"),
        today,
    )
    accounts = load_recent_douyin_accounts(
        raw_file,
        wps["raw_data_sheet"],
        assignments,
        lookback_start,
        today,
    )

    entries = tuple(
        RosterEntry(
            anchor_code=assignment.anchor_code,
            streamer_name=assignment.streamer_name,
            makeup_artist=assignment.makeup_artist or "化妆师未填写",
            phone=assignment.phone,
            douyin_accounts=tuple(sorted(accounts[anchor_code])),
            team=assignment.team,
        )
        for anchor_code, assignment in sorted(assignments.items())
        if anchor_code in accounts
    )
    account_index: dict[str, RosterEntry] = {}
    for entry in entries:
        for account in entry.douyin_accounts:
            previous = account_index.get(account)
            if previous and previous.anchor_code != entry.anchor_code:
                raise ValueError(
                    f"抖音号 {account} 同时对应主播编号 "
                    f"{previous.anchor_code} 和 {entry.anchor_code}。"
                )
            account_index[account] = entry

    LOGGER.info(
        "名单刷新完成: 当前有效主播 %d 个，其中手机号非空 %d 个、小队已匹配 %d 个；"
        "近 %d 天抖音号 %d 个；名单=%s；原始数据=%s",
        len(entries),
        sum(bool(entry.phone) for entry in entries),
        sum(bool(entry.team) for entry in entries),
        lookback_days,
        len(account_index),
        roster_file.name,
        raw_file.name,
    )
    return RosterSnapshot(
        roster_file=roster_file,
        raw_file=raw_file,
        entries=entries,
        account_index=account_index,
        lookback_start=lookback_start,
        lookback_end=today,
    )
